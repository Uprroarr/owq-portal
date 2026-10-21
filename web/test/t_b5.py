"""v121 batch 5: financial statements (owner only), they reconcile, GAAP / IFRS labels, the .xlsx opens as a real workbook."""
import sys,os,time,json,subprocess
SP='/tmp/claude-0/-home-claude/6e37954b-9c9e-57d7-b212-24d756de5469/scratchpad';sys.path.insert(0,SP)
import lgx
from playwright.sync_api import sync_playwright
PORT=8865;SITE=os.environ.get('SITE',SP+'/web/test/site_b5');URL='http://127.0.0.1:%d/index.html'%PORT
res=[]
def ok(c,m,info=''):res.append(bool(c));print(('PASS ' if c else 'FAIL ')+m+('' if c else '  '+str(info)[:500]),flush=True)
srv=subprocess.Popen(['python3','-m','http.server',str(PORT),'--bind','127.0.0.1'],cwd=SITE,stdout=subprocess.DEVNULL,stderr=subprocess.DEVNULL);time.sleep(1)
try:
 with sync_playwright() as p:
  b=p.chromium.launch(executable_path='/opt/pw-browsers/chromium',args=['--no-sandbox','--disable-webgl','--no-proxy-server'])
  ctx=b.new_context(viewport={'width':1440,'height':900},accept_downloads=True);ctx.add_init_script("try{localStorage.setItem('owq_gq','still')}catch(e){}")
  errs=[]
  K=ctx.new_page();K.goto('http://127.0.0.1:%d/keep.html'%PORT);K.wait_for_function('window.__FAKE')
  A=ctx.new_page();A.on('pageerror',lambda e:errs.append(str(e)[:200]))
  A.add_init_script("window.__FAKE_SIGNIN=%s;"%json.dumps({'uid':'uOwner','email':'owner@example.com','displayName':'Cole Leckey'}))
  A.goto(URL);A.wait_for_function("window.__FAKE",timeout=20000);A.evaluate("__FAKE.call('reset',{owner:'owner@example.com'})");A.evaluate("sessionStorage.clear()");A.reload()
  A.wait_for_function("OWQC.phase==='out'",timeout=30000);A.click('#owqgi');A.wait_for_function("OWQC.phase==='in'",timeout=30000)
  lgx.login(A,2);A.wait_for_function("typeof SYN!=='undefined'&&SYN.first",timeout=40000);A.evaluate("(()=>{try{mock()}catch(e){}})()");A.wait_for_timeout(1200)
  A.evaluate("openTab('Agency Performance','Capital')");A.wait_for_timeout(1200)
  A.evaluate("CAP.rows.push({d:dkey(new Date()),t:'Private equity',who:'Fund',amt:50000,al:'Business',dir:'in'},{d:dkey(new Date()),t:'Loan',who:'Bank',amt:20000,al:'Business',dir:'in'},{d:dkey(new Date()),t:'Owner distribution',amt:5000,al:'Personal',dir:'out'});capSave()")
  A.evaluate("openTab('Agency Performance','Financial Statements')");A.wait_for_timeout(1500)
  t=A.inner_text('#main').upper()
  ok('INCOME STATEMENT' in t and 'STATEMENT OF CASH FLOWS' in t and 'BALANCE SHEET' in t and 'NOTES' in t,'income statement, cash flows, balance sheet and notes',t[:300])
  F=A.evaluate("fsData()")
  ok(abs((F['cash']+F['inv'])-(F['debt']+F['own']+F['eqv']+F['re']+F['openCash']))<0.5,'the balance sheet balances',F)
  ok(abs(F['ni']-(F['rev']-F['opx']+F['real']))<0.01 and F['fin']['eq']==50000 and F['fin']['debt']==20000 and F['fin']['dist']==5000,'net income and financing flows add up',F['fin'])
  lead=A.evaluate("(()=>{const [a,b]=fsRange();return sum(D.expenses.filter(x=>isLeadEx(x)&&x.d>=a&&x.d<=b),x=>+x.amt||0)})()")
  la=[r for r in F['opex'] if r[0]=='Lead acquisition'];ok(la and abs(la[0][1]-lead)<1,'lead buys are the lead acquisition line',[la,lead])
  A.click(".fsbar .seg button:has-text('IFRS')");A.wait_for_timeout(500);t=A.inner_text('#main').upper()
  ok('STATEMENT OF PROFIT OR LOSS' in t and 'STATEMENT OF FINANCIAL POSITION' in t and 'PROFIT FOR THE PERIOD' in t,'IFRS presentation uses IFRS names')
  ok('NOT CERTIFIED' in t and 'CPA' in t,'clear note that they are unaudited management statements')
  A.screenshot(path=SP+'/web/test/shots/b5_fs.png')
  with A.expect_download() as dl:A.click(".fsact .btn:has-text('Google Sheets (.xlsx)')")
  path=dl.value.path();import shutil;shutil.copy(path,SP+'/web/test/fs_test.xlsx')
  import openpyxl;wb=openpyxl.load_workbook(SP+'/web/test/fs_test.xlsx')
  ok(len(wb.sheetnames)==4 and any('Profit' in n or 'Income' in n for n in wb.sheetnames),'the .xlsx opens as a real workbook with a tab per statement',wb.sheetnames)
  ws=wb[wb.sheetnames[0]];vals=[c.value for row in ws.iter_rows() for c in row if isinstance(c.value,(int,float))]
  ok(len(vals)>3 and any(abs(v-round(F['ni'],2))<0.02 for v in vals),'numbers come through as numbers',vals[:6])
  with A.expect_download() as dl2:A.click(".fsact .btn:has-text('CSV')")
  ok(dl2.value.suggested_filename.endswith('.csv'),'CSV export too')
  # agents cannot open it
  ok('Financial Statements' in A.evaluate("apxTabs().join(',')"),'owner has the tab')
  A.evaluate("WHO='John Montini';openTab('Agency Performance','Financial Statements')");A.wait_for_timeout(500)
  ok(A.evaluate("SUB['Agency Performance']")=='Overview' and 'INCOME STATEMENT' not in A.inner_text('#main').upper(),'agents cannot open the financial statements')
  ok(not errs,'no page errors',errs)
  b.close()
finally:srv.terminate()
print('%d/%d'%(sum(res),len(res)))
